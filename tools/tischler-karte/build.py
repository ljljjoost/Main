"""Baut die Tischler-Karte aus der Ranking-Tabelle.

Aufruf: python3 build.py <Tischler-Praktikum_Hamburg_Ranking.xlsx> [fragment.html]

Schreibt docs/tischler/index.html (vollständige Seite für GitHub Pages).
Mit zweitem Argument zusätzlich eine Fassung ohne <html>/<head>/<body>-Gerüst (für ein Claude-Artifact).
basemap.json muss vorher mit basemap.py erzeugt worden sein.
"""
import json, os, re, sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "..", "docs", "tischler", "index.html")

# Spaltenüberschrift (ohne Zeilenumbruch) -> Schlüssel im Seiten-JSON
COLS = {
    "Rang": "r", "Betrieb": "n", "Inhaltl. Fit (0–10)": "fit", "Entfernung Winterhude (km)": "km",
    "Praktikum-Status": "st", "Chance (0–3)": "ch", "Typ": "typ", "Ähnlich wie": "aw", "Stadtteil": "ort",
    "Schwerpunkt": "sp", "Einschätzung": "ein", "Praktikum – Details": "det", "Telefon": "tel",
    "E-Mail": "mail", "Website / Quelle-Link": "web", "Adresse": "adr", "Restaurierung (0–3)": "c0",
    "Entwurf & Möbelbau (0–3)": "c1", "Feinarbeit & Gestaltung (0–2)": "c2", "Hand + Maschine (0–2)": "c3",
    "Quelle(n)": "src", "Breite (OSM)": "lat", "Länge (OSM)": "lng",
}


def read_rows(path):
    ws_f = openpyxl.load_workbook(path)["Ranking"]               # Hyperlinks
    ws_v = openpyxl.load_workbook(path, data_only=True)["Ranking"]  # berechnete Werte
    head = [re.sub(r"\s*\n\s*", " ", str(c.value or "")).strip() for c in ws_v[1]]
    idx = {COLS[h]: i for i, h in enumerate(head) if h in COLS}
    missing = set(COLS.values()) - set(idx)
    if missing:
        sys.exit(f"Spalten fehlen in der Tabelle: {sorted(missing)}")
    rows = []
    for vals, cells in zip(ws_v.iter_rows(min_row=2, values_only=True), ws_f.iter_rows(min_row=2)):
        if not vals[idx["n"]]:
            continue
        d = {k: vals[i] for k, i in idx.items()}
        code, _, label = str(d["st"]).partition("·")
        d["st"], d["stl"] = int(code), label.strip()
        link = cells[idx["web"]].hyperlink
        d["href"] = link.target if link else (("https://" + d["web"]) if d["web"] else None)
        d["c"] = [d.pop(f"c{i}") for i in range(4)]
        if d["aw"] in (None, "–", "-"):
            d["aw"] = None
        rows.append({k: v for k, v in d.items() if v is not None})
    rows.sort(key=lambda d: d["r"])
    return rows


def read_stand(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Legende"]
    for row in ws.iter_rows(max_row=5, values_only=True):
        for v in row:
            m = re.search(r"Stand:\s*([\d.]+)", str(v or ""))
            if m:
                return m.group(1)
    return "–"


def main():
    rows = read_rows(sys.argv[1])
    tpl = open(os.path.join(HERE, "template.html"), encoding="utf-8").read()
    css = open(os.path.join(HERE, "vendor", "leaflet-1.9.4.css"), encoding="utf-8").read()
    base = open(os.path.join(HERE, "basemap.json"), encoding="utf-8").read()
    js = lambda s: s.replace("</", "<\\/")
    page = (tpl.replace("/*LEAFLET_CSS*/", css)
               .replace("__STAND__", read_stand(sys.argv[1]))
               .replace("__DATA__", js(json.dumps(rows, ensure_ascii=False, separators=(",", ":"))))
               .replace("__BASEMAP__", js(base)))
    head, body = page.split("<!--/HEAD-->")
    head = head.replace("<!--HEAD-->", "").strip()
    full = ("<!doctype html>\n<html lang=\"de\">\n<head>\n<meta charset=\"utf-8\">\n"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1, viewport-fit=cover\">\n"
            f"{head}\n</head>\n<body>\n{body.strip()}\n</body>\n</html>\n")
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    open(OUT, "w", encoding="utf-8").write(full)
    print(f"{len(rows)} Betriebe -> {os.path.normpath(OUT)} ({len(full) // 1024} KB)")
    if len(sys.argv) > 2:
        open(sys.argv[2], "w", encoding="utf-8").write(head + "\n" + body.strip() + "\n")
        print("Fragment ->", sys.argv[2])


if __name__ == "__main__":
    main()
